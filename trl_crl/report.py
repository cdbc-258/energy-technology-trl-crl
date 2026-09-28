"""Current-method/current-result reports only."""
from pathlib import Path
import json

def write_report(dest,outputs):
 s=outputs['SUMMARY']
 text=f'''# 技术TRL/CRL评估结果

评价资料截止：{s['assessment_cutoff']}。

|项目|数量|
|---|---:|
|文献主题|{s['theme_count']}|
|候选技术方向|{s['candidate_direction_count']}|
|有界评价对象|{s['bounded_object_count']}|
|TRL有值|{s['objects_with_trl']}|
|CRL有值|{s['objects_with_crl']}|
|同对象两轴都有值|{s['objects_with_both']}|
|两轴均未知|{s['objects_with_neither']}|

## 当前方法

{s['theme_count']}主题目录用于组织文献；识别技术名称和路线后，限定配置、功能、场景与时期，形成评价对象。只有与对象一致的实际验证和商业事件可用作证据。

TRL依据Q/GDW 12566—2025的九级定义，检查技术状态、集成状态与试验环境。CRL依据招标书五级定义及调研细化的操作规则。两轴分别核算；专利申请或授权不独立决定CRL。逐级所需门槛全部获得支持才输出该证据阶段，有冲突或缺证据的门槛不计通过。

当前采用{s['accepted_evidence_passages']}段证据引文、{s['reviewed_gate_observations']}条门槛观测，来源URL计数为{s['source_url_count']}，另列{s['observed_fact_count']}条工程/商业量化事实。实际结果、额定容量、年化估算和未来目标分别标注。

## 如何读结果

`assessment_units.csv`或`技术_TRL_CRL评估结果.xlsx`的“技术对象结果”列出全部对象、两轴和边界。空值表示证据不足或未完成所需判据，不代表0级或1级。输出阶段是已有完整证据支持的阶段，不是实际成熟度的绝对上限。

C开头与W开头是采集来源形成的对象标识，均按同一套规则评审；W编号用于网页/项目案例标识，不是等级。一个案例可能关联多个方向，不能跨方向重复计数为多个独立技术。

{s['candidate_direction_count']}个方向以关联案例的证据分布展示，方向统一等级为空。来源可能为论文、政府汇编、项目报告或供应商自述，并非全部独立核验。证据时期可能是2019、2020等，不能默认等同于截止日产业现状。

## 使用限制

所有结果为公开证据初评，正式等级均为空。关键技术风险清单、专属细则审定、自评与独立专家程序尚不完整。招标书36项指标的完整采集、专家权重和阈值标定未完成，不输出加权综合分或招标验收结论。部分对象仍缺测试合格要求或商业履约材料，详见“证据缺项”工作表。

公开仓库保留必要引文、URL和引文哈希，不再分发标准/招标全文和完整论文。离线验证检查引文内部完整性和判据实现，不宣称重新验证全部远端源文或证明语义准确率。
'''
 (Path(dest)/'评估报告.md').write_text(text,encoding='utf-8')

def write_workbook(dest,d,o):
 from openpyxl import Workbook
 from openpyxl.styles import Font,PatternFill,Alignment
 from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
 from openpyxl.utils import get_column_letter
 wb=Workbook();wb.remove(wb.active)
 def sheet(title,rows,cols):
  ws=wb.create_sheet(title);ws.append([label for key,label in cols]);ws.freeze_panes='A2'
  for row in rows:
   vals=[]
   for key,label in cols:
    v=row.get(key)
    if isinstance(v,(list,dict)):v=json.dumps(v,ensure_ascii=False)
    if isinstance(v,str):
     v=ILLEGAL_CHARACTERS_RE.sub('',v)
     if len(v)>32000:v=v[:31800]+' [完整内容见JSON]'
     if v.startswith(('=','+','-','@')):v="'"+v
    vals.append(v)
   ws.append(vals)
  ws.auto_filter.ref=ws.dimensions
  for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17365D')
  for i,(key,label) in enumerate(cols,1):ws.column_dimensions[get_column_letter(i)].width=42 if key in {'quote','rationale','boundary','required_evidence','remaining_gaps','interpretation'} else 25
  for row in ws.iter_rows(min_row=2):
   for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
   ws.row_dimensions[row[0].row].height=48
 sheet('当前统计',[dict(item=k,value=v) for k,v in o['SUMMARY'].items()],[('item','指标'),('value','数值/口径')])
 sheet('技术对象结果',o['assessment_units'],[('case_id','对象ID'),('canonical_name','技术对象'),('trl_public_evidence_stage','TRL证据阶段'),('crl_public_evidence_stage','CRL证据阶段'),('paired_axes_usable','同对象双轴齐全'),('evidence_period','证据时期'),('object_configuration','配置'),('application_or_target_function','应用场景'),('boundary','边界'),('remaining_gaps','待补证据'),('formal_trl','正式TRL'),('formal_crl','正式CRL')])
 sheet('技术方向证据分布',o['direction_evidence_profiles'],[('technology_id','方向ID'),('canonical_name','技术方向'),('case_relations','案例与关系'),('associated_case_trl_distribution','关联案例TRL分布'),('associated_case_crl_distribution','关联案例CRL分布'),('same_case_pairs','同对象双轴'),('direction_trl','方向TRL'),('direction_crl','方向CRL'),('interpretation','解释')])
 sheet('主题目录',d['themes'],[('category_id','主题ID'),('category_name','主题名称'),('documents','记录数'),('candidate_direction_names','候选技术方向'),('processing_status','识别状态'),('processing_reason','说明')])
 sheet('证据缺项',o['evidence_gaps'],[('case_id','对象ID'),('canonical_name','技术对象'),('missing_axes','未知轴'),('evidence_period','证据时期'),('required_evidence','所需证据')])
 sheet('逐项判据',[g for a in o['assessments'] for g in a['gates']],[('case_id','对象ID'),('axis','评估轴'),('level','阶段'),('key','判据键'),('requirement','要求'),('effective_status','状态'),('rationale','理由'),('observation_ids','观测引用')])
 sheet('来源目录',d['sources'],[('source_id','来源ID'),('title','题名'),('url','URL'),('published_at','公开日期'),('date_note','日期说明'),('quality_note','来源性质')])
 sheet('证据引文',d['evidence'],[('evidence_id','证据ID'),('case_id','对象ID'),('source_url','来源URL'),('quote','原文引文'),('quote_start','原文字符起点'),('quote_end','原文字符终点'),('acceptance_scope','适用边界'),('quote_sha256','引文SHA256')])
 sheet('量化事实',d['observed_facts'],[('case_id','对象ID'),('metric','指标'),('value','数值'),('unit','单位'),('value_type','数据类型'),('note','限定'),('source_url','URL'),('quote','原文')])
 reqs=[dict(case_id=p['case_id'],**r) for p in d['technical_profiles'] for r in p['performance_requirements']]
 sheet('性能要求',reqs,[('case_id','对象ID'),('requirement_id','要求ID'),('metric','指标'),('target','目标'),('unit','单位'),('basis','依据状态'),('basis_quote','依据引文'),('critical_for_stage','适用阶段'),('note','说明')])
 wb.save(Path(dest)/'技术_TRL_CRL评估结果.xlsx')
