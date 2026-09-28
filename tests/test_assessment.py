import copy,json,tempfile,unittest
from pathlib import Path
from trl_crl.pipeline import assess,load_inputs,write_outputs
from trl_crl.engine import evaluate_case
ROOT=Path(__file__).resolve().parents[1]

class AssessmentTests(unittest.TestCase):
 @classmethod
 def setUpClass(cls):cls.data,cls.outputs=assess(ROOT/'data')
 def sample(self):
  d=self.data;c=next(x for x in d['objects'] if x['case_id']=='W001');r=next(x for x in d['gate_reviews'] if x['case_id']=='W001');p=next(x for x in d['technical_profiles'] if x['case_id']=='W001');o=[x for x in d['observations'] if x['case_id']=='W001'];e={x['evidence_id']:x for x in d['evidence']}
  return copy.deepcopy((c,r,p,o,e))
 def test_current_counts_and_unique_ids(self):
  s=self.outputs['SUMMARY'];self.assertEqual((s['bounded_object_count'],s['objects_with_trl'],s['objects_with_crl'],s['objects_with_both']),(122,108,23,13))
  self.assertEqual(len({u['case_id'] for u in self.outputs['assessment_units']}),122)
 def test_axes_are_independent(self):
  c,r,p,o,e=self.sample()
  for st in r['axes']['CRL']:
   st['review_state']='unreviewed'
   for g in st['gates']:g.update(status='unknown',observation_ids=[])
  a=evaluate_case(c,r,o,e,p,'2026-09-25')['axes'];self.assertEqual(a['TRL']['public_evidence_stage'],9);self.assertIsNone(a['CRL']['public_evidence_stage'])
 def test_reject_cross_object_observation(self):
  c,r,p,o,e=self.sample();o[0]['case_id']='W013'
  with self.assertRaises(ValueError):evaluate_case(c,r,o,e,p,'2026-09-25')
 def test_reject_planned_event_marked_complete(self):
  c,r,p,o,e=self.sample();o[0]['assertion']='planned'
  with self.assertRaises(ValueError):evaluate_case(c,r,o,e,p,'2026-09-25')
 def test_unresolved_contradiction_blocks_stage(self):
  c,r,p,o,e=self.sample();negative=copy.deepcopy(o[0]);negative['observation_id']='negative';negative['polarity']='contradict';o.append(negative)
  a=evaluate_case(c,r,o,e,p,'2026-09-25');self.assertIsNone(a['axes']['TRL']['public_evidence_stage'])
 def test_criteria_change_requires_review(self):
  c,r,p,o,e=self.sample();p['version']=100
  with self.assertRaises(ValueError):evaluate_case(c,r,o,e,p,'2026-09-25')
 def test_no_cross_configuration_hydrogen_pair(self):
  m={u['case_id']:u for u in self.outputs['assessment_units']}
  self.assertEqual(m['C0191-B01']['trl_public_evidence_stage'],6);self.assertIsNone(m['C0191-B01']['crl_public_evidence_stage'])
  self.assertIsNone(m['C0191-B02']['trl_public_evidence_stage']);self.assertEqual(m['C0191-B02']['crl_public_evidence_stage'],2)
 def test_synthetic_income_and_future_capacity_not_rated(self):
  m={u['case_id']:u for u in self.outputs['assessment_units']}
  self.assertIsNone(m['C0524-N01']['crl_public_evidence_stage']);self.assertIsNone(m['W013']['trl_public_evidence_stage'])
  self.assertTrue(any(f['case_id']=='C0524-N01' and f['value_type']=='synthetic_price_calculation' for f in self.data['observed_facts']))
 def test_integrity_validator_rejects_changed_quote(self):
  import shutil
  with tempfile.TemporaryDirectory() as tmp:
   t=Path(tmp)/'data';shutil.copytree(ROOT/'data',t);p=t/'evidence.json';e=json.loads(p.read_text());e[0]['quote']='tampered';p.write_text(json.dumps(e))
   with self.assertRaises(ValueError):load_inputs(t)
 def test_reject_future_evidence(self):
  import shutil
  with tempfile.TemporaryDirectory() as tmp:
   t=Path(tmp)/'data';shutil.copytree(ROOT/'data',t);p=t/'evidence.json';e=json.loads(p.read_text());e[0]['published_at']='2099-01-01';p.write_text(json.dumps(e))
   with self.assertRaises(ValueError):load_inputs(t)
 def test_portable_export(self):
  from openpyxl import load_workbook
  with tempfile.TemporaryDirectory() as tmp:
   write_outputs(ROOT/'data',tmp);wb=load_workbook(Path(tmp)/'技术_TRL_CRL评估结果.xlsx',read_only=True)
   self.assertEqual(wb['技术对象结果'].max_row,123);self.assertEqual(wb['技术方向证据分布'].max_row,392)
   self.assertFalse(any('对照' in s or '前后' in s for s in wb.sheetnames));wb.close()
if __name__=='__main__':unittest.main()
